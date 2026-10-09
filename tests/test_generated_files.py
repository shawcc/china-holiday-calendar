import csv
import json
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


class GeneratedFileTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.manifest = json.loads((ROOT / "public" / "manifest.json").read_text(encoding="utf-8"))
        cls.latest = cls.manifest["years"][-1]

    def test_manifest_counts_match_source(self):
        holidays = self.latest["holidays"]
        rest = sum(len(item["rest_dates"]) for item in holidays)
        work = sum(len(item["work_dates"]) for item in holidays)
        self.assertEqual(self.latest["rest_count"], rest)
        self.assertEqual(self.latest["work_count"], work)
        self.assertEqual(self.latest["row_count"], rest + work)

    def test_xlsx_matches_import_contract(self):
        workbook = load_workbook(ROOT / "public" / self.latest["xlsx"])
        sheet = workbook["工作日历规则导入模板"]
        rows = [
            (sheet.cell(row, 1).value, sheet.cell(row, 2).value)
            for row in range(3, 3 + self.latest["row_count"])
        ]
        self.assertEqual(len(rows), 39)
        self.assertEqual(rows[0], ("2026-01-01", "休息日"))
        self.assertEqual(rows[-1], ("2026-10-10", "工作日"))
        self.assertTrue(all(day_type in {"工作日", "休息日"} for _, day_type in rows))
        self.assertEqual(len(sheet.data_validations.dataValidation), 1)

    def test_csv_matches_xlsx_rows(self):
        with (ROOT / "public" / self.latest["csv"]).open(encoding="utf-8-sig") as file:
            csv_rows = list(csv.reader(file))[1:]
        workbook = load_workbook(ROOT / "public" / self.latest["xlsx"], read_only=True)
        sheet = workbook["工作日历规则导入模板"]
        xlsx_rows = [
            [sheet.cell(row, 1).value, sheet.cell(row, 2).value]
            for row in range(3, 3 + self.latest["row_count"])
        ]
        self.assertEqual(csv_rows, xlsx_rows)


if __name__ == "__main__":
    unittest.main()
