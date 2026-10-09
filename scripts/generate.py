#!/usr/bin/env python3
import csv
import hashlib
import json
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
OUTPUT_DIR = ROOT / "public" / "downloads"


def load_year(path: Path) -> dict:
    with path.open(encoding="utf-8") as file:
        data = json.load(file)
    if data["status"] != "verified":
        raise ValueError(f"{path.name} is not verified")
    return data


def rows_for(data: dict) -> list[tuple[str, str, str]]:
    rows = []
    for holiday in data["holidays"]:
        rows.extend((date, "休息日", holiday["name"]) for date in holiday["rest_dates"])
        rows.extend((date, "工作日", holiday["name"]) for date in holiday["work_dates"])
    rows.sort(key=lambda row: row[0])
    dates = [row[0] for row in rows]
    if len(dates) != len(set(dates)):
        raise ValueError(f"{data['year']} contains duplicate dates")
    return rows


def build_xlsx(data: dict, rows: list[tuple[str, str, str]], target: Path) -> None:
    workbook = Workbook()
    release_time = datetime.fromisoformat(data["published_at"])
    workbook.properties.creator = "China Holiday Calendar"
    workbook.properties.created = release_time
    workbook.properties.modified = release_time
    sheet = workbook.active
    sheet.title = "工作日历规则导入模板"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "填写说明：请从第3行开始填写；日期格式为 YYYY-MM-DD；类型只能选择“工作日”或“休息日”。"
    sheet["A1"].font = Font(name="宋体", size=11, bold=True, color="9C5700")
    sheet["A1"].fill = PatternFill("solid", fgColor="FFEB9C")
    sheet["A1"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 46

    for column, value in (("A2", "日期"), ("B2", "类型")):
        sheet[column] = value
        sheet[column].font = Font(name="宋体", size=11, bold=True, color="FFFFFF")
        sheet[column].fill = PatternFill("solid", fgColor="1F4E78")

    for index, (date, day_type, _) in enumerate(rows, start=3):
        sheet.cell(index, 1, date).number_format = "@"
        sheet.cell(index, 2, day_type)
        sheet.cell(index, 1).font = Font(name="宋体", size=11)
        sheet.cell(index, 2).font = Font(name="宋体", size=11)

    validation = DataValidation(type="list", formula1='"工作日,休息日"')
    sheet.add_data_validation(validation)
    validation.add("B3:B1000")
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 41.18
    for index in range(2, 1001):
        sheet.row_dimensions[index].height = 17

    workbook.save(target)
    normalize_xlsx(target, release_time)


def normalize_xlsx(path: Path, release_time: datetime) -> None:
    normalized = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                root = ET.fromstring(content)
                namespace = "{http://purl.org/dc/terms/}"
                timestamp = release_time.isoformat(timespec="seconds") + "Z"
                root.find(f"{namespace}created").text = timestamp
                root.find(f"{namespace}modified").text = timestamp
                content = ET.tostring(root, encoding="utf-8", xml_declaration=False)
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = source.getinfo(name).external_attr
            target.writestr(info, content)
    normalized.replace(path)


def build_csv(rows: list[tuple[str, str, str]], target: Path) -> None:
    with target.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["日期", "类型"])
        writer.writerows((date, day_type) for date, day_type, _ in rows)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    manifest = {"latest_year": None, "years": []}
    for path in sorted(DATA_DIR.glob("*.json")):
        data = load_year(path)
        rows = rows_for(data)
        year = data["year"]
        xlsx_name = f"中国节假日工作日历-{year}.xlsx"
        csv_name = f"中国节假日工作日历-{year}.csv"
        xlsx_path = OUTPUT_DIR / xlsx_name
        csv_path = OUTPUT_DIR / csv_name
        build_xlsx(data, rows, xlsx_path)
        build_csv(rows, csv_path)
        rest_count = sum(day_type == "休息日" for _, day_type, _ in rows)
        work_count = sum(day_type == "工作日" for _, day_type, _ in rows)
        manifest["years"].append(
            {
                **data,
                "row_count": len(rows),
                "rest_count": rest_count,
                "work_count": work_count,
                "xlsx": f"downloads/{xlsx_name}",
                "csv": f"downloads/{csv_name}",
                "xlsx_sha256": sha256(xlsx_path),
                "csv_sha256": sha256(csv_path)
            }
        )
        manifest["latest_year"] = max(manifest["latest_year"] or year, year)

    (ROOT / "public" / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8"
    )


if __name__ == "__main__":
    main()
